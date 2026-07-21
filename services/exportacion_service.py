"""
services/exportacion_service.py   ──   RF-40 Exportar reportes
════════════════════════════════════════════════════════════════

HU: Como Operador de Bienestar, quiero exportar CUALQUIER reporte a
Excel o PDF con el formato oficial de la UNAB (logo institucional,
nombre de la institución, fecha de generación y usuario que lo
generó) para usarlo en informes, auditorías y presentaciones
directivas.

¿QUÉ HACE?
  Exporta cualquier reporte (una lista de dicts + definición de
  columnas) a Excel, PDF o CSV, con el formato oficial de la UNAB.
  El logo y el nombre de la institución se cargan dinámicamente desde
  ConfiguracionService (el mismo que usa el login y el resto del
  sistema) — si no hay logo configurado, el documento se genera
  igual, solo que sin la imagen.

¿POR QUÉ UN SERVICIO SEPARADO DE ReporteService?
  ReporteService arma los DATOS (RF-36 a RF-39, RF-41). Este servicio
  solo sabe convertir "columnas + filas" en un archivo — no le importa
  si esos datos vienen de aptos, bienes, inscritos, etc. Así cualquier
  reporte nuevo se exporta gratis, sin tocar este archivo.

USO TÍPICO (desde la UI, tras generar un reporte):
    columnas = [
        ("codigo", "Código"),
        ("nombre", "Nombre"),
        ("taller_nombre", "Taller"),
        ("porcentaje_asistencia", "% Asistencia"),
    ]
    resultado = ExportacionService.exportar_excel(
        titulo="Reporte de aptos para certificado",
        columnas=columnas,
        filas=reporte.lista,
        usuario_nombre=sesion.nombre_completo,
        ruta_destino="/mnt/user-data/outputs/aptos_certificado.xlsx",
    )
    if resultado.ok:
        abrir_archivo(resultado.ruta)
    else:
        QMessageBox.warning(self, "No se pudo exportar", resultado.mensaje)

RENDIMIENTO (< 5 s):
  Cada export es una sola pasada secuencial sobre `filas` (sin
  consultas a BD adicionales — los datos ya vienen calculados desde
  ReporteService) más, como mucho, una carga de imagen del logo desde
  disco. Para los volúmenes típicos de un reporte de talleres
  (decenas a un par de miles de filas) esto corre muy por debajo de
  5 segundos tanto en Excel (openpyxl) como en PDF (reportlab).

DEPENDENCIAS OPCIONALES:
  • exportar_excel() requiere `openpyxl`      (pip install openpyxl)
  • exportar_pdf()   requiere `reportlab`      (pip install reportlab)
  • exportar_csv()   no requiere nada (usa el módulo csv de la stdlib)
  Si falta la librería, el método retorna ok=False con un mensaje
  claro en vez de lanzar una excepción no controlada.
"""

import csv
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any, Tuple

from services.configuracion_service import ConfiguracionService


NOMBRE_INSTITUCION_DEFECTO = "UNAB — Sistema de Talleres Extracurriculares"

# Definición de columna: (clave_en_el_dict, etiqueta_visible)
Columna = Tuple[str, str]


@dataclass
class ResultadoExportacion:
    """Estructura de respuesta estándar de este servicio."""
    ok:      bool
    mensaje: str
    ruta:    Optional[str] = None


class ExportacionService:
    """RF-40: Exportar cualquier reporte a Excel, PDF o CSV, con el
    formato oficial de la UNAB (logo, institución, fecha, usuario)."""

    # ══════════════════════════════════════════════════════════════
    # Datos institucionales (logo + nombre) — mismo origen que el
    # login y el resto del sistema (configuracion.json).
    # ══════════════════════════════════════════════════════════════
    @staticmethod
    def _obtener_datos_institucionales() -> Tuple[str, Optional[str]]:
        """
        Retorna (nombre_institucion, ruta_absoluta_logo_o_None).
        Nunca lanza excepción: si algo falla, cae al nombre por
        defecto y sin logo, para que la exportación nunca se rompa
        por un problema de configuración.
        """
        nombre = NOMBRE_INSTITUCION_DEFECTO
        ruta_logo = None
        try:
            resultado = ConfiguracionService.obtener()
            if resultado.ok:
                config = resultado.datos
                nombre = config.get("nombre_institucion") or nombre

                logo_ruta = config.get("logo_ruta")
                if logo_ruta:
                    ruta_absoluta = ConfiguracionService.obtener_ruta_logo_absoluta(logo_ruta)
                    if ruta_absoluta and Path(ruta_absoluta).exists():
                        ruta_logo = ruta_absoluta
        except Exception:
            pass
        return nombre, ruta_logo

    # ══════════════════════════════════════════════════════════════
    # CSV
    # ══════════════════════════════════════════════════════════════
    @staticmethod
    def exportar_csv(
        titulo: str,
        columnas: List[Columna],
        filas: List[Dict[str, Any]],
        usuario_nombre: str,
        ruta_destino: str,
    ) -> ResultadoExportacion:
        try:
            nombre_institucion, _ = ExportacionService._obtener_datos_institucionales()

            ruta = Path(ruta_destino)
            ruta.parent.mkdir(parents=True, exist_ok=True)

            # utf-8-sig para que Excel reconozca tildes/ñ al abrir el CSV.
            with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow([nombre_institucion])
                writer.writerow([titulo])
                writer.writerow([
                    f"Generado por: {usuario_nombre}",
                    f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
                ])
                writer.writerow([])
                writer.writerow([etiqueta for _, etiqueta in columnas])
                for fila in filas:
                    writer.writerow([fila.get(clave, "") for clave, _ in columnas])

            return ResultadoExportacion(
                ok=True,
                mensaje=f"Reporte exportado correctamente a {ruta.name}.",
                ruta=str(ruta),
            )
        except Exception as e:
            return ResultadoExportacion(ok=False, mensaje=f"Error al exportar CSV: {e}")

    # ══════════════════════════════════════════════════════════════
    # EXCEL — con logo institucional real (igual que ExportarExcel.lista_aptos)
    # ══════════════════════════════════════════════════════════════
    @staticmethod
    def exportar_excel(
        titulo: str,
        columnas: List[Columna],
        filas: List[Dict[str, Any]],
        usuario_nombre: str,
        ruta_destino: str,
    ) -> ResultadoExportacion:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill
            from openpyxl.utils import get_column_letter
            from openpyxl.drawing.image import Image as ImagenExcel
        except ImportError:
            return ResultadoExportacion(
                ok=False,
                mensaje="Falta instalar 'openpyxl' para exportar a Excel "
                        "(pip install openpyxl).",
            )

        try:
            nombre_institucion, ruta_logo = ExportacionService._obtener_datos_institucionales()

            ruta = Path(ruta_destino)
            ruta.parent.mkdir(parents=True, exist_ok=True)

            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte"
            total_columnas = max(len(columnas), 4)

            # ── Logo (si hay) + encabezado institucional ────────────
            # Con logo: el texto arranca en la columna C para no
            # superponerse con la imagen (igual patrón que
            # ExportarExcel.lista_aptos, que usa "C2" para el título).
            if ruta_logo:
                try:
                    logo = ImagenExcel(ruta_logo)
                    logo.width = 64
                    logo.height = 64
                    ws.add_image(logo, "A1")
                    ws.row_dimensions[1].height = 48
                    ws.row_dimensions[2].height = 22
                    col_inicio = 3
                except Exception:
                    col_inicio = 1
            else:
                col_inicio = 1

            ws.merge_cells(start_row=1, start_column=col_inicio, end_row=1, end_column=total_columnas)
            ws.cell(row=1, column=col_inicio, value=nombre_institucion).font = Font(bold=True, size=13)

            ws.merge_cells(start_row=2, start_column=col_inicio, end_row=2, end_column=total_columnas)
            ws.cell(row=2, column=col_inicio, value=titulo).font = Font(bold=True, size=11)

            ws.cell(row=3, column=col_inicio, value=f"Generado por: {usuario_nombre}")
            ws.cell(row=3, column=col_inicio + 2, value=f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}")

            # ── Encabezado de la tabla ────────────────────────────
            fila_encabezado = 5
            for col, (_, etiqueta) in enumerate(columnas, start=1):
                celda = ws.cell(row=fila_encabezado, column=col, value=etiqueta)
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="4F46E5")
                celda.alignment = Alignment(horizontal="center")

            for i, fila_datos in enumerate(filas, start=fila_encabezado + 1):
                for col, (clave, _) in enumerate(columnas, start=1):
                    ws.cell(row=i, column=col, value=fila_datos.get(clave, ""))

            for col in range(1, len(columnas) + 1):
                ws.column_dimensions[get_column_letter(col)].width = 22

            wb.save(ruta)

            return ResultadoExportacion(
                ok=True,
                mensaje=f"Reporte exportado correctamente a {ruta.name}.",
                ruta=str(ruta),
            )
        except Exception as e:
            return ResultadoExportacion(ok=False, mensaje=f"Error al exportar Excel: {e}")

    # ══════════════════════════════════════════════════════════════
    # PDF — con logo institucional real (igual que ExportarPDF.lista_aptos)
    # ══════════════════════════════════════════════════════════════
    @staticmethod
    def exportar_pdf(
        titulo: str,
        columnas: List[Columna],
        filas: List[Dict[str, Any]],
        usuario_nombre: str,
        ruta_destino: str,
    ) -> ResultadoExportacion:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.platypus import (
                SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
                Image as ImagenPDF
            )
        except ImportError:
            return ResultadoExportacion(
                ok=False,
                mensaje="Falta instalar 'reportlab' para exportar a PDF "
                        "(pip install reportlab).",
            )

        try:
            nombre_institucion, ruta_logo = ExportacionService._obtener_datos_institucionales()

            ruta = Path(ruta_destino)
            ruta.parent.mkdir(parents=True, exist_ok=True)

            doc = SimpleDocTemplate(
                str(ruta),
                pagesize=landscape(A4),
                topMargin=1.2 * cm,
                bottomMargin=1.5 * cm,
                leftMargin=1.2 * cm,
                rightMargin=1.2 * cm,
            )
            estilos = getSampleStyleSheet()

            estilo_institucion = ParagraphStyle(
                "Institucion", parent=estilos["Heading2"],
                fontSize=13, textColor=colors.HexColor("#1E3A8A"), spaceAfter=2,
            )
            estilo_titulo_reporte = ParagraphStyle(
                "TituloReporte", parent=estilos["Heading3"],
                fontSize=11, textColor=colors.HexColor("#374151"), spaceAfter=2,
            )
            estilo_meta = ParagraphStyle(
                "Meta", parent=estilos["Normal"],
                fontSize=9, textColor=colors.HexColor("#6B7280"),
            )

            # ── Encabezado institucional: logo (si hay) + textos ────
            celda_textos = [
                Paragraph(nombre_institucion, estilo_institucion),
                Paragraph(titulo, estilo_titulo_reporte),
                Paragraph(
                    f"Generado por: {usuario_nombre}  |  "
                    f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                    estilo_meta,
                ),
            ]

            if ruta_logo:
                try:
                    logo_flowable = ImagenPDF(ruta_logo, width=1.8 * cm, height=1.8 * cm)
                except Exception:
                    logo_flowable = Paragraph("", estilos["Normal"])
            else:
                logo_flowable = Paragraph("", estilos["Normal"])

            tabla_encabezado = Table(
                [[logo_flowable, celda_textos]],
                colWidths=[2.4 * cm, None]
            )
            tabla_encabezado.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (0, 0), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))

            elementos = [tabla_encabezado, Spacer(1, 14)]

            # ── Tabla de datos ───────────────────────────────────────
            encabezados = [etiqueta for _, etiqueta in columnas]
            datos_tabla = [encabezados]
            for fila_datos in filas:
                datos_tabla.append([str(fila_datos.get(clave, "")) for clave, _ in columnas])

            tabla = Table(datos_tabla, repeatRows=1)
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            elementos.append(tabla)

            doc.build(elementos)

            return ResultadoExportacion(
                ok=True,
                mensaje=f"Reporte exportado correctamente a {ruta.name}.",
                ruta=str(ruta),
            )
        except Exception as e:
            return ResultadoExportacion(ok=False, mensaje=f"Error al exportar PDF: {e}")