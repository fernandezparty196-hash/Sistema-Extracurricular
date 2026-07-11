"""
Utilidades para importar inscripciones masivas desde un archivo Excel.
No accede a la base de datos: solo lee/escribe archivos .xlsx.
"""
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def leer_plantilla_inscripcion(ruta_archivo: str) -> list[dict]:
    """
    Lee un archivo Excel con la plantilla de inscripción masiva.

    Formato esperado (fila 1 = encabezados):
        Columna 'DNI'    -> obligatoria
        Columna 'Nombre' -> opcional, solo referencial (no se usa para validar)

    Retorna una lista de dicts:
        [{"fila": 2, "dni": "12345678", "nombre_ref": "..."}, ...]

    Lanza ValueError si el archivo no tiene el formato esperado.
    """
    try:
        wb = load_workbook(ruta_archivo, data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"No se pudo abrir el archivo Excel: {e}")

    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    if not filas:
        raise ValueError("El archivo está vacío.")

    encabezados = [str(c).strip().upper() if c else "" for c in filas[0]]

    if "DNI" not in encabezados:
        raise ValueError(
            "El archivo debe tener una columna llamada 'DNI' en la primera fila."
        )

    idx_dni = encabezados.index("DNI")
    idx_nombre = encabezados.index("NOMBRE") if "NOMBRE" in encabezados else None

    registros = []
    for i, fila in enumerate(filas[1:], start=2):
        if fila is None or all(c is None for c in fila):
            continue  # fila vacía, se ignora

        dni_raw = fila[idx_dni] if idx_dni < len(fila) else None
        nombre_raw = (
            fila[idx_nombre]
            if idx_nombre is not None and idx_nombre < len(fila)
            else None
        )

        dni = str(dni_raw).strip() if dni_raw is not None else ""
        # Excel a veces guarda el DNI como número (float) y quita ceros a la izquierda
        if dni.endswith(".0"):
            dni = dni[:-2]
        if dni.isdigit() and len(dni) < 8:
            dni = dni.zfill(8)

        registros.append({
            "fila": i,
            "dni": dni,
            "nombre_ref": str(nombre_raw).strip() if nombre_raw else "",
        })

    if not registros:
        raise ValueError("El archivo no contiene filas de datos.")

    return registros


def generar_plantilla_inscripcion(ruta_destino: str) -> None:
    """Genera un archivo Excel de plantilla vacío para inscripción masiva."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inscripción masiva"

    encabezados = ["DNI", "Nombre (referencial)"]
    ws.append(encabezados)

    header_fill = PatternFill(start_color="1D9E75", end_color="1D9E75", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=col)
        celda.fill = header_fill
        celda.font = header_font
        celda.alignment = Alignment(horizontal="center")

    # Fila de ejemplo para guiar al usuario
    ws.append(["71234567", "Ejemplo — puedes borrar esta fila"])

    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 38

    wb.save(ruta_destino)
